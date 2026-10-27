#!/usr/bin/env python3
"""Print a JSON snapshot of page, typography, paragraph, and table formatting."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError as exc:
    raise SystemExit("python-docx is required in the active Python environment") from exc


def pt(value):
    return None if value is None else round(value / 12700, 2)


def cm(value):
    return None if value is None else round(value / 360000, 2)


def run_info(run):
    rpr = run._r.rPr
    rfonts = rpr.rFonts if rpr is not None else None
    color = run.font.color
    return {
        "text": run.text[:80],
        "size_pt": pt(run.font.size),
        "bold": run.bold,
        "italic": run.italic,
        "color": str(color.rgb) if color and color.rgb else None,
        "ascii": rfonts.get(qn("w:ascii")) if rfonts is not None else None,
        "east_asia": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
    }


def paragraph_info(paragraph):
    fmt = paragraph.paragraph_format
    return {
        "text": paragraph.text[:120],
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": str(fmt.alignment) if fmt.alignment is not None else None,
        "line_spacing": fmt.line_spacing,
        "space_before_pt": pt(fmt.space_before),
        "space_after_pt": pt(fmt.space_after),
        "first_line_indent_cm": cm(fmt.first_line_indent),
        "left_indent_cm": cm(fmt.left_indent),
        "runs": [run_info(run) for run in paragraph.runs[:4]],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if not args.docx.exists():
        raise SystemExit(f"DOCX not found: {args.docx}")

    doc = Document(str(args.docx))
    report = {"docx": str(args.docx.resolve()), "sections": [], "styles": {}, "samples": []}

    for index, section in enumerate(doc.sections):
        report["sections"].append(
            {
                "index": index,
                "page_width_cm": cm(section.page_width),
                "page_height_cm": cm(section.page_height),
                "top_margin_cm": cm(section.top_margin),
                "bottom_margin_cm": cm(section.bottom_margin),
                "left_margin_cm": cm(section.left_margin),
                "right_margin_cm": cm(section.right_margin),
                "header_distance_cm": cm(section.header_distance),
                "footer_distance_cm": cm(section.footer_distance),
                "start_type": str(section.start_type),
            }
        )

    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        try:
            style = doc.styles[name]
            fmt = style.paragraph_format
            report["styles"][name] = {
                "font_name": style.font.name,
                "font_size_pt": pt(style.font.size),
                "bold": style.font.bold,
                "color": str(style.font.color.rgb)
                if style.font.color and style.font.color.rgb
                else None,
                "alignment": str(fmt.alignment) if fmt.alignment is not None else None,
                "line_spacing": fmt.line_spacing,
                "space_before_pt": pt(fmt.space_before),
                "space_after_pt": pt(fmt.space_after),
                "first_line_indent_cm": cm(fmt.first_line_indent),
            }
        except KeyError:
            pass

    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip():
            item = paragraph_info(paragraph)
            item["index"] = index
            report["samples"].append(item)
        if len(report["samples"]) >= 18:
            break

    for table_index, table in enumerate(doc.tables):
        if len(table.rows) > 1 and len(table.columns) > 1:
            report["table_sample"] = {
                "table_index": table_index,
                "rows": len(table.rows),
                "columns": len(table.columns),
                "header": paragraph_info(table.cell(0, 0).paragraphs[0]),
                "body": paragraph_info(table.cell(1, 0).paragraphs[0]),
            }
            break

    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(payload, encoding="utf-8")
        print(f"wrote {args.output}")
    else:
        print(payload)


if __name__ == "__main__":
    sys.exit(main())
