#!/usr/bin/env python3
"""Audit a sourced Chinese popular-science DOCX.

Checks citation/reference wiring, duplicate markers, figure/table numbering,
image and page counts, hyperlink targets, and basic caption source notes.
It does not verify scientific truth; pair it with the editorial checklist.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError as exc:
    raise SystemExit("python-docx is required in the active Python environment") from exc

CITE_RE = re.compile(r"\[\d+(?:[-,]\d+)*\]")
REF_RE = re.compile(r"^\[(\d+)\]")
FIG_RE = re.compile(r"^图\s*(\d+)\b")
TABLE_RE = re.compile(r"^表\s*(\d+)\b")
DUP_RE = re.compile(r"(\[\d+(?:[-,]\d+)*\])\1+")


def collect_paragraphs(doc):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        for part in (section.header, section.footer):
            paragraphs.extend(part.paragraphs)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
    return paragraphs


def expand_citation(token):
    values = set()
    for part in token.strip("[]").split(","):
        if "-" in part:
            start, end = map(int, part.split("-", 1))
            values.update(range(min(start, end), max(start, end) + 1))
        elif part.isdigit():
            values.add(int(part))
    return values


def inspect_media(path):
    groups = defaultdict(list)
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        for name in archive.namelist():
            if name.startswith("word/media/") and not name.endswith("/"):
                groups[hashlib.sha256(archive.read(name)).hexdigest()].append(name)
    return bad_member, groups


def sequential_issues(kind, numbers):
    numbers = sorted(set(numbers))
    expected = list(range(1, len(numbers) + 1))
    if numbers != expected:
        return f"{kind} numbering is not sequential: {numbers}"
    return None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--expect-images", type=int)
    parser.add_argument("--expect-figures", type=int)
    parser.add_argument("--expect-tables", type=int, help="number of numbered table captions")
    parser.add_argument("--expect-pages", type=int)
    parser.add_argument("--required-phrase", action="append", default=[])
    parser.add_argument("--forbidden-phrase", action="append", default=[])
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not args.docx.exists():
        raise SystemExit(f"DOCX not found: {args.docx}")

    errors = []
    warnings = []
    report = {"docx": str(args.docx.resolve())}

    try:
        doc = Document(str(args.docx))
    except Exception as exc:
        raise SystemExit(f"cannot open DOCX: {exc}") from exc

    paragraphs = collect_paragraphs(doc)
    text_blob = "\n".join(p.text for p in paragraphs)

    reference_numbers = {}
    in_text_numbers = set()
    duplicate_markers = []
    figures = []
    tables = []

    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        duplicate_markers.extend(m.group(0) for m in DUP_RE.finditer(text))

        m = REF_RE.match(text)
        if m:
            number = int(m.group(1))
            reference_numbers[number] = text
            continue

        for token in CITE_RE.findall(text):
            in_text_numbers.update(expand_citation(token))

        m = FIG_RE.match(text)
        if m:
            figures.append((int(m.group(1)), text))
        m = TABLE_RE.match(text)
        if m:
            tables.append((int(m.group(1)), text))

    if duplicate_markers:
        errors.append(f"duplicated citation markers: {duplicate_markers[:10]}")

    if not reference_numbers:
        warnings.append("no numbered reference entries found")

    missing_references = sorted(in_text_numbers - set(reference_numbers))
    uncited_references = sorted(set(reference_numbers) - in_text_numbers)
    if missing_references:
        errors.append(f"citations without reference entries: {missing_references}")
    if uncited_references:
        errors.append(f"reference entries never cited: {uncited_references}")

    bookmark_names = {
        node.get(qn("w:name"))
        for node in doc._element.xpath(".//w:bookmarkStart")
        if node.get(qn("w:name"))
    }
    anchors = [
        node.get(qn("w:anchor"))
        for node in doc._element.xpath(".//w:hyperlink")
        if node.get(qn("w:anchor"))
    ]
    missing_anchors = sorted(set(anchors) - bookmark_names)
    if missing_anchors:
        errors.append(f"internal links without bookmark anchors: {missing_anchors}")

    expected_ref_bookmarks = {f"ref{n}" for n in reference_numbers}
    missing_ref_bookmarks = sorted(expected_ref_bookmarks - bookmark_names)
    if missing_ref_bookmarks:
        errors.append(f"missing reference bookmarks: {missing_ref_bookmarks}")

    external_targets = []
    bad_external_targets = []
    external_relationships = [rel for rel in doc.part.rels.values() if rel.is_external]
    for rel in external_relationships:
        target = rel.target_ref
        if not target.startswith(("http://", "https://")):
            bad_external_targets.append(target)
    for hyperlink in doc._element.xpath(".//w:hyperlink"):
        rid = hyperlink.get(qn("r:id"))
        if rid and rid in doc.part.rels:
            target = doc.part.rels[rid].target_ref
            external_targets.append(target)
    if bad_external_targets:
        errors.append(f"non-http external targets: {bad_external_targets}")

    figure_numbers = [n for n, _ in figures]
    table_numbers = [n for n, _ in tables]
    for kind, numbers in (("figure", figure_numbers), ("table", table_numbers)):
        issue = sequential_issues(kind, numbers)
        if issue:
            errors.append(issue)

    for number, caption in figures:
        if not any(marker in caption for marker in ("来源", "依据", "整理自", "[")):
            warnings.append(f"figure {number} caption may lack a source note: {caption[:100]}")

    try:
        bad_member, media_groups = inspect_media(args.docx)
    except zipfile.BadZipFile as exc:
        raise SystemExit(f"invalid DOCX ZIP package: {exc}") from exc
    if bad_member:
        errors.append(f"corrupt ZIP member: {bad_member}")

    report.update(
        {
            "text_characters": len(text_blob),
            "inline_images": len(doc.inline_shapes),
            "document_tables": len(doc.tables),
            "numbered_figures": figure_numbers,
            "numbered_tables": table_numbers,
            "references": len(reference_numbers),
            "cited_reference_numbers": sorted(in_text_numbers),
            "missing_references": missing_references,
            "uncited_references": uncited_references,
            "bookmark_count": len(bookmark_names),
            "reference_bookmarks": sorted(expected_ref_bookmarks & bookmark_names),
            "internal_link_count": len(anchors),
            "external_link_count": len(external_targets),
            "external_relationship_count": len(external_relationships),
            "media_total": sum(len(v) for v in media_groups.values()),
            "media_unique": len(media_groups),
        }
    )

    if report["media_total"] > len(doc.inline_shapes):
        warnings.append(
            f"DOCX contains {report['media_total']} media parts but {len(doc.inline_shapes)} inline images; unused media may remain in the package"
        )

    if args.expect_images is not None and len(doc.inline_shapes) != args.expect_images:
        errors.append(
            f"expected {args.expect_images} inline images, found {len(doc.inline_shapes)}"
        )
    if args.expect_figures is not None and len(set(figure_numbers)) != args.expect_figures:
        errors.append(
            f"expected {args.expect_figures} numbered figures, found {len(set(figure_numbers))}"
        )
    if args.expect_tables is not None and len(set(table_numbers)) != args.expect_tables:
        errors.append(
            f"expected {args.expect_tables} numbered tables, found {len(set(table_numbers))}"
        )

    for phrase in args.required_phrase:
        if phrase not in text_blob:
            errors.append(f"required phrase missing: {phrase!r}")
    for phrase in args.forbidden_phrase:
        if phrase in text_blob:
            errors.append(f"forbidden phrase present: {phrase!r}")

    if args.pdf:
        if not args.pdf.exists():
            errors.append(f"PDF not found: {args.pdf}")
        else:
            try:
                import pypdfium2 as pdfium
            except ImportError:
                warnings.append("pypdfium2 unavailable; PDF page count not checked")
            else:
                pdf = pdfium.PdfDocument(str(args.pdf))
                report["pdf_pages"] = len(pdf)
                if args.expect_pages is not None and len(pdf) != args.expect_pages:
                    errors.append(
                        f"expected {args.expect_pages} PDF pages, found {len(pdf)}"
                    )

    report["errors"] = errors
    report["warnings"] = warnings
    report["status"] = "PASS" if not errors else "FAIL"

    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"STATUS {report['status']}")
        print(f"DOCX {args.docx}")
        print(f"REFERENCES {report['references']}")
        print(f"CITED {report['cited_reference_numbers']}")
        print(f"FIGURES {report['numbered_figures']}")
        print(f"TABLES {report['numbered_tables']}")
        print(f"IMAGES inline={report['inline_images']} media={report['media_total']} unique={report['media_unique']}")
        print(f"LINKS internal={report['internal_link_count']} external={report['external_link_count']}")
        if "pdf_pages" in report:
            print(f"PDF_PAGES {report['pdf_pages']}")
        for warning in warnings:
            print(f"WARNING {warning}")
        for error in errors:
            print(f"ERROR {error}")

    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
